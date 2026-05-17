"""
AI Supply Chain Platform — Client Presentation Builder
Professional white-background consulting-style deck.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Colour palette (light / professional) ─────────────────────────────────────
NAVY        = RGBColor(0x0B, 0x1F, 0x3B)   # headings, text
TEAL        = RGBColor(0x00, 0x9E, 0x96)   # primary accent
TEAL_LIGHT  = RGBColor(0xE0, 0xF5, 0xF4)   # teal tint for backgrounds
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE   = RGBColor(0xF8, 0xFA, 0xFC)   # slide background
LIGHT_GREY  = RGBColor(0xE9, 0xEE, 0xF4)   # card backgrounds
MID_GREY    = RGBColor(0x6B, 0x7A, 0x90)   # sub-text
DARK_GREY   = RGBColor(0x2D, 0x3A, 0x4A)   # body text
RED         = RGBColor(0xDC, 0x26, 0x26)
RED_LIGHT   = RGBColor(0xFE, 0xF2, 0xF2)   # red tint
AMBER       = RGBColor(0xD9, 0x77, 0x06)
AMBER_LIGHT = RGBColor(0xFF, 0xF7, 0xED)
GREEN       = RGBColor(0x05, 0x96, 0x69)
GREEN_LIGHT = RGBColor(0xEC, 0xFD, 0xF5)   # green tint
PURPLE      = RGBColor(0x70, 0x3A, 0xD4)
PURPLE_LIGHT= RGBColor(0xF5, 0xF0, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ══════════════════════════════════════════════════════════════════════════════
# HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def fill_bg(slide, color=OFF_WHITE):
    add_rect(slide, 0, 0, 13.33, 7.5, fill=color)

def add_rect(slide, x, y, w, h, fill=None, line=None, line_width=Pt(1), radius=False):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, x, y, w, h, size=11, bold=False, color=DARK_GREY,
             align=PP_ALIGN.LEFT, italic=False):
    txb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txb

def h_rule(slide, y, x=0.4, w=12.53, color=LIGHT_GREY, thickness=Pt(1)):
    add_rect(slide, x, y, w, 0.015, fill=color)

def accent_bar(slide, y=0, h=7.5, x=0, w=0.22, color=TEAL):
    """Vertical left accent stripe."""
    add_rect(slide, x, y, w, h, fill=color)

def top_band(slide, color=NAVY, height=1.15):
    add_rect(slide, 0, 0, 13.33, height, fill=color)

def footer(slide, left_text="AI Supply Chain Intelligence Platform",
           right_text=""):
    add_rect(slide, 0, 7.18, 13.33, 0.32, fill=LIGHT_GREY)
    add_text(slide, left_text,  0.35, 7.2,  9,    0.28, size=7.5, color=MID_GREY)
    if right_text:
        add_text(slide, right_text, 10.5, 7.2, 2.5, 0.28, size=7.5,
                 color=MID_GREY, align=PP_ALIGN.RIGHT)

def kpi_badge(slide, x, y, label, value, accent=TEAL):
    """Small metric badge: accent top bar + white card."""
    add_rect(slide, x, y,       1.85, 0.1,  fill=accent)
    add_rect(slide, x, y+0.1,   1.85, 0.75, fill=WHITE,
             line=LIGHT_GREY, line_width=Pt(1))
    add_text(slide, value, x+0.1, y+0.15, 1.65, 0.35,
             size=15, bold=True, color=accent, align=PP_ALIGN.CENTER)
    add_text(slide, label, x+0.1, y+0.5,  1.65, 0.28,
             size=7.5, color=MID_GREY, align=PP_ALIGN.CENTER)

def flow_panel(slide, x, y, w, h, steps, accent):
    """3-step How It Works flow panel — replaces screenshot placeholder."""
    # Background
    add_rect(slide, x, y, w, h, fill=NAVY)
    # Header
    add_rect(slide, x, y, w, 0.36, fill=RGBColor(0x06, 0x14, 0x28))
    add_text(slide, "  HOW IT WORKS", x+0.1, y+0.07, w-0.2, 0.24,
             size=8, bold=True, color=accent)

    step_h   = 0.82
    arrow_h  = 0.2
    sy = y + 0.46
    for i, (heading, detail) in enumerate([(h, d) for h, d in [(s[0], s[1]) for s in steps]]):
        # Step card
        add_rect(slide, x+0.15, sy, w-0.3, step_h,
                 fill=RGBColor(0x0F, 0x2D, 0x52))
        add_rect(slide, x+0.15, sy, 0.06, step_h, fill=accent)
        # Step number bubble
        add_rect(slide, x+0.28, sy+0.13, 0.28, 0.28, fill=accent)
        add_text(slide, str(i+1), x+0.28, sy+0.13, 0.28, 0.28,
                 size=9, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        # Heading
        add_text(slide, heading, x+0.65, sy+0.08, w-0.88, 0.26,
                 size=8.5, bold=True, color=WHITE)
        # Detail
        add_text(slide, detail, x+0.28, sy+0.44, w-0.5, 0.34,
                 size=7.5, color=RGBColor(0xB0, 0xC4, 0xD8))
        sy += step_h
        # Arrow between steps
        if i < 2:
            add_text(slide, "\u25bc", x + w/2 - 0.12, sy+0.01, 0.28, 0.18,
                     size=8, color=accent, align=PP_ALIGN.CENTER)
            sy += arrow_h


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — COVER
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK)
fill_bg(slide, OFF_WHITE)

# Left navy panel
add_rect(slide, 0, 0, 6.5, 7.5, fill=NAVY)
# Teal bottom accent on left panel
add_rect(slide, 0, 6.3, 6.5, 1.2, fill=TEAL)

# Product label
add_text(slide, "AI-POWERED SUPPLY CHAIN INTELLIGENCE",
         0.45, 0.5, 5.7, 0.35, size=8.5, bold=True,
         color=RGBColor(0xAA,0xCC,0xDD), align=PP_ALIGN.LEFT)

# Main title
add_text(slide,
         "Transforming Supply\nChain Decision-Making\nwith Real-Time AI",
         0.45, 1.05, 5.7, 2.8, size=34, bold=True, color=WHITE)

# Subtitle
add_text(slide,
         "8 Real-World Scenarios — From Problem to\nResolution in Minutes, Not Days.",
         0.45, 4.0, 5.5, 0.9, size=13, color=RGBColor(0xCC,0xE8,0xF0))

# Teal rule
add_rect(slide, 0.45, 3.8, 3.5, 0.04, fill=TEAL)

# URL on teal band
add_text(slide, "baf-unified.vercel.app",
         0.45, 6.45, 5.5, 0.35, size=10, bold=True, color=NAVY)
add_text(slide, "Live Platform  ·  Demo Ready",
         0.45, 6.8, 5.5, 0.28, size=8.5, color=NAVY)

# Right panel — 3 layer cards
add_text(slide, "THREE LAYERS. ONE PLATFORM.",
         6.85, 0.55, 6.1, 0.32, size=8.5, bold=True, color=MID_GREY)
add_text(slide, "Complete Visibility Across\nYour Entire Supply Chain",
         6.85, 0.88, 6.1, 1.0, size=20, bold=True, color=NAVY)
h_rule(slide, 1.95, x=6.85, w=6.05, color=TEAL, thickness=Pt(2))

layer_cards = [
    (TEAL,   "WDE — Warehouse Diagnosis Engine",
     "Real-time KPI monitoring, AI exception detection,\nauto-generated corrective actions."),
    (AMBER,  "TMS — Supply Master Transportation",
     "Live shipment tracking, carrier performance,\ncost & delay analytics across all modes."),
    (GREEN,  "Planning — Demand & Supply Hub",
     "Demand-supply gap analysis, stockout early\nwarning, Gen AI assistant for instant queries."),
]
for i, (color, title, desc) in enumerate(layer_cards):
    by = 2.15 + i * 1.6
    add_rect(slide, 6.75, by,      0.07, 1.3, fill=color)
    add_rect(slide, 6.82, by,      6.15, 1.3, fill=WHITE,
             line=LIGHT_GREY, line_width=Pt(1))
    add_text(slide, title, 7.0, by+0.1,  5.8, 0.32, size=11, bold=True, color=NAVY)
    add_text(slide, desc,  7.0, by+0.48, 5.8, 0.65, size=9.5, color=DARK_GREY)

# Bottom right tag
add_rect(slide, 6.75, 7.0, 6.22, 0.38, fill=LIGHT_GREY)
add_text(slide, "Single Login  ·  Voice-Enabled (VAPI)  ·  AI Backend on Railway  ·  Vercel Deployed",
         6.85, 7.04, 6.0, 0.28, size=8, color=MID_GREY, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — PLATFORM OVERVIEW
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK)
fill_bg(slide, OFF_WHITE)
top_band(slide, NAVY, 1.05)
accent_bar(slide, color=TEAL)

add_text(slide, "PLATFORM OVERVIEW",
         0.45, 0.15, 12, 0.3, size=8.5, bold=True, color=TEAL)
add_text(slide, "One Platform. Three Layers. Full Supply Chain Intelligence.",
         0.45, 0.48, 12, 0.48, size=20, bold=True, color=WHITE)

# 3 detailed cards
detail_cards = [
    (TEAL,  "WDE", "Warehouse Diagnosis Engine", [
        "Real-time KPI monitoring (pick rate, accuracy, utilisation)",
        "AI drift monitor — detects decline before it becomes a crisis",
        "Exception Action Panel with auto-generated corrective actions",
        "One-click 'Approve & Send to Ops Team' — no email chains",
        "Inbound / outbound exception classification by severity",
    ]),
    (AMBER, "TMS", "Supply Master Transportation", [
        "Live tracking: on-time rate, delay rate, cost per shipment",
        "Carrier & route-level drill-down for instant root cause",
        "Filters by carrier, mode, region, date range, order type",
        "CO\u2082 emissions tracking for ESG & sustainability reporting",
        "Tender acceptance rate — negotiation intelligence",
    ]),
    (GREEN, "Planning", "Demand & Supply Planning Hub", [
        "Upload CSV — demand vs supply gap analysis in seconds",
        "Stockout early warning weeks before peak season",
        "Month-on-month trend across SKUs, warehouses, regions",
        "Gen AI Assistant — query incident history in plain English",
        "Data Hub: upload data to all 3 layers from one screen",
    ]),
]

for i, (accent, short, title, points) in enumerate(detail_cards):
    bx = 0.35 + i * 4.3
    # Card
    add_rect(slide, bx, 1.25, 4.05, 5.55, fill=WHITE,
             line=LIGHT_GREY, line_width=Pt(1))
    # Top colour header
    add_rect(slide, bx, 1.25, 4.05, 0.55, fill=accent)
    add_text(slide, short,  bx+0.15, 1.28, 0.6,  0.25, size=8,  bold=True, color=NAVY)
    add_text(slide, title,  bx+0.15, 1.5,  3.75, 0.28, size=10, bold=True, color=NAVY)
    # Bullet points
    y = 1.95
    for pt in points:
        add_rect(slide, bx+0.2, y+0.12, 0.07, 0.07, fill=accent)
        add_text(slide, pt, bx+0.38, y, 3.5, 0.42, size=9, color=DARK_GREY)
        y += 0.46

footer(slide, right_text="Platform Overview")


# ══════════════════════════════════════════════════════════════════════════════
# SCENARIO DATA
# ══════════════════════════════════════════════════════════════════════════════
scenarios = [
    {
        "num": 1,
        "layer": "Planning — Demand & Supply",
        "accent": GREEN,
        "accent_light": GREEN_LIGHT,
        "title": "Peak Season Demand Surge — Beverage SKUs at Stockout Risk",
        "trigger": "Summer or festive season approaches. Demand for Food & Beverage SKUs accelerates faster than the supply plan. The gap is only discovered when shelves are already short.",
        "without": [
            "Demand plan and supply plan managed in separate Excel files across teams",
            "No single view of which SKUs are at risk — analysis is manual and slow",
            "Stockout discovered too late — reorder window has already passed",
            "Retailers face empty shelves during the highest-revenue period of the year",
        ],
        "with_tool": [
            "Platform shows Stockout Risk (High SKUs) and Days of Cover in real time",
            "Demand Variance KPI flags SKUs where actual demand outpaces the plan",
            "Filter by Category: Food & Beverages — isolate the relevant portfolio instantly",
            "Supply team acts early — inventory secured before the demand peak hits",
        ],
        "outcome": "Stockout risk surfaced early across the beverage portfolio. Supply team repositions inventory before retailer shelves are impacted.",
        "capability": "Stockout Risk & Days of Cover — Planning Dashboard",
        "ai_feature": "Email Alert: When Stockout Risk crosses the threshold, an automated alert is sent to the Supply Planning head — flagging exactly which SKUs need immediate action.",
        "steps": [
            ("Demand Signal",  "Platform detects demand outpacing supply for key SKUs"),
            ("KPI Fires Red",  "Stockout Risk KPI and Days of Cover drop below threshold"),
            ("Team Acts Early","Supply team secures inventory before the demand peak hits"),
        ],
    },
    {
        "num": 2,
        "layer": "WDE — Warehouse Layer",
        "accent": TEAL,
        "accent_light": TEAL_LIGHT,
        "title": "Inbound Receiving Bottleneck Slowing the Entire Warehouse",
        "trigger": "Inbound trucks are arriving on schedule but goods are sitting at the dock for hours before being put away. Dock-to-Stock time is climbing and nobody has flagged it.",
        "without": [
            "Dock-to-Stock time tracked manually — no live visibility",
            "Receiving accuracy issues only caught after stock is already misplaced",
            "Warehouse manager finds out from a shift supervisor — too late in the day",
            "Outbound dispatch gets impacted as inbound congestion builds up",
        ],
        "with_tool": [
            "WDE Inbound Velocity page tracks Dock-to-Stock time live across shifts",
            "Receiving Accuracy KPI flags mismatches before stock enters the system",
            "Disruption Risk Banner alerts the manager the moment a threshold is crossed",
            "Exception Action Panel generates corrective actions — sent to Ops in one click",
        ],
        "outcome": "Inbound bottleneck caught within the shift, not at day-end. Corrective action reaches the right team before outbound dispatch is affected.",
        "capability": "Inbound Velocity & Receiving Accuracy — WDE Dashboard",
        "ai_feature": "VAPI Calling Feature: When Dock-to-Stock time crosses the threshold, the AI agent automatically calls the Warehouse Manager, asks for a justification, and records the response as a transcript.",
        "steps": [
            ("Threshold Breach",  "Dock-to-Stock time crosses limit — Disruption Risk Banner fires"),
            ("AI Generates Plan", "Exception Action Panel creates ranked corrective actions instantly"),
            ("One-Click Dispatch","Ops team receives the action plan — bottleneck cleared same shift"),
        ],
    },
    {
        "num": 3,
        "layer": "Unified Dashboard — Cross-Layer",
        "accent": PURPLE,
        "accent_light": PURPLE_LIGHT,
        "title": "Perfect Order Rate Declining — Retailer Fill Rate Under Pressure",
        "trigger": "Retail partners start flagging incomplete orders. The ops team is looking at warehouse, logistics, and planning data separately — nobody sees the full picture.",
        "without": [
            "Perfect Order Rate not tracked as a single metric — it lives in three systems",
            "Warehouse blames logistics, logistics blames planning — no shared KPI",
            "Retailer complaints arrive before the internal team has identified the issue",
            "Resolution requires a cross-functional meeting that takes days to convene",
        ],
        "with_tool": [
            "Unified Dashboard shows Perfect Order Rate as a single cross-layer KPI",
            "Drills into Order Fill Rate (WMS), On-Time Delivery (TMS), Demand Variance (Planning)",
            "Supply Chain Reliability score shows which layer is pulling the metric down",
            "All three teams see the same number — aligned instantly, not in 3 days",
        ],
        "outcome": "Perfect Order Rate tracked in one place. Root layer identified immediately — all three teams aligned and acting the same day.",
        "capability": "Perfect Order Rate & Supply Chain Reliability — Unified Dashboard",
        "ai_feature": "Gen AI Assistant: When Perfect Order Rate dips, any team member can ask 'What should we do now?' — the assistant draws on past incident knowledge and recommends the right corrective action.",
        "steps": [
            ("Single KPI View",  "Perfect Order Rate shown as one cross-layer metric on the dashboard"),
            ("Layer Drill-Down", "Platform identifies which layer — WMS, TMS, or Planning — is failing"),
            ("Same-Day Fix",     "All three teams aligned on the same number — resolution coordinated instantly"),
        ],
    },
    {
        "num": 4,
        "layer": "TMS — Transportation Layer",
        "accent": AMBER,
        "accent_light": AMBER_LIGHT,
        "title": "Carrier Underperforming on a Key Distribution Lane",
        "trigger": "A carrier handling a critical distribution lane shows a steady decline in on-time performance. The ops team has no visibility until a distributor escalates.",
        "without": [
            "On-time data lives in the carrier's own portal — no consolidated view",
            "Decline is gradual — no threshold alert fires until it becomes critical",
            "Tender acceptance rate not tracked — carrier negotiation done blind",
            "By the time it is escalated, multiple distributors are already affected",
        ],
        "with_tool": [
            "TMS Carrier Performance page tracks On-Time Rate per carrier in real time",
            "Platform compares current period vs previous period — decline flagged automatically",
            "Tender Acceptance Rate gives negotiation leverage before the next contract cycle",
            "Delay Rate by route pinpoints worst-hit lanes — filter in one click",
        ],
        "outcome": "Carrier performance decline caught early, before distributors are impacted. Procurement negotiates backed by live data, not anecdote.",
        "capability": "Carrier On-Time Rate & Tender Acceptance — TMS",
        "ai_feature": "VAPI Calling Feature: When On-Time Rate drops below threshold, the AI agent calls the Logistics Head, asks for a justification, and records the conversation as a transcript for review.",
        "steps": [
            ("Auto Comparison",  "Platform compares current vs previous period for every carrier"),
            ("Decline Flagged",  "Falling on-time rate and tender acceptance rate surfaced immediately"),
            ("Data-Led Action",  "Procurement renegotiates or reallocates — backed by live lane data"),
        ],
    },
    {
        "num": 5,
        "layer": "WDE — Warehouse Layer",
        "accent": TEAL,
        "accent_light": TEAL_LIGHT,
        "title": "Dispatch Backlog Building — Orders Not Moving Out on Time",
        "trigger": "Order pendency is climbing. Outbound dispatches are falling behind plan. The floor team is aware but the escalation has not been triggered.",
        "without": [
            "Order Pendency visible only in the next morning's daily report",
            "On-Time Dispatch KPI not monitored intra-day — shift ends before anyone acts",
            "Backlog compounds across shifts — by Day 3 it is a full crisis",
            "Customer SLAs breached before leadership is even aware of the issue",
        ],
        "with_tool": [
            "WDE Dispatch Backlog page shows pending orders live across all shifts",
            "Order Pendency and On-Time Dispatch KPIs update continuously",
            "Threshold breach triggers an exception — AI generates a prioritised action plan",
            "Ops team receives the plan instantly — backlog cleared before it escalates",
        ],
        "outcome": "Dispatch backlog flagged within the same shift. Ops team mobilised before the backlog compounds or a single customer SLA is missed.",
        "capability": "Dispatch Backlog & Order Pendency — WDE Live KPIs",
        "ai_feature": "Email Alert: When Order Pendency crosses the threshold, an automated alert is sent to the Operations Manager — with the current backlog count and the breached KPI highlighted.",
        "steps": [
            ("Live Detection",   "Order Pendency KPI crosses threshold — exception raised mid-shift"),
            ("AI Action Plan",   "AI generates a prioritised backlog recovery plan automatically"),
            ("Instant Dispatch", "Ops team acts on the plan — backlog cleared before day-end"),
        ],
    },
    {
        "num": 6,
        "layer": "Resilient Control Tower — All Layers",
        "accent": PURPLE,
        "accent_light": PURPLE_LIGHT,
        "title": "Supply Chain Disruption — See It Before the Customer Does",
        "trigger": "A supplier delay hits inbound. The warehouse gets congested. Outbound slows. Carriers start missing windows. Each team sees only their own piece of the problem.",
        "without": [
            "Supplier, warehouse, logistics, and planning data in four separate systems",
            "No end-to-end view — disruption propagates silently across the chain",
            "Leadership gets a fragmented picture across multiple status calls",
            "Customer feels the disruption before any team has a complete view",
        ],
        "with_tool": [
            "Resilient Control Tower: Supplier \u2192 Warehouse \u2192 Transport \u2192 Customer on one screen",
            "Each node shows a live health score — green, amber, or red at a glance",
            "Flow arrows between nodes show where the bottleneck sits right now",
            "Cross-layer alerts surface automatically — no status calls needed",
        ],
        "outcome": "End-to-end supply chain health on one screen. Disruption pinpointed to the exact node — response coordinated before the customer is impacted.",
        "capability": "End-to-End Node Health — Resilient Control Tower",
        "ai_feature": "VAPI Calling Feature: When the Control Tower detects a cross-layer disruption, the AI agent calls the Supply Chain Head, briefs them on the affected nodes, and records their response as an accountable transcript.",
        "steps": [
            ("Node Turns Red",   "Disruption detected at one node — health score drops to red"),
            ("Chain Impact Mapped","AI correlates impact across all downstream nodes automatically"),
            ("Coordinated Fix",  "Response team sees the full picture — acts before customer is affected"),
        ],
    },
    {
        "num": 7,
        "layer": "TMS — Transportation Layer",
        "accent": AMBER,
        "accent_light": AMBER_LIGHT,
        "title": "Freight Cost Visibility Without Waiting for Finance",
        "trigger": "Freight costs on key distribution lanes are rising. The logistics team suspects it but cannot quantify it without waiting for the monthly finance consolidation.",
        "without": [
            "Cost per shipment tracked only after carrier invoices are reconciled",
            "Route-level visibility requires manual extraction from multiple sources",
            "Finance consolidation takes weeks — every decision is retrospective",
            "Procurement negotiates the next contract without current cost data",
        ],
        "with_tool": [
            "TMS Cost Analytics shows Cost per Shipment live — no finance wait",
            "Route-level breakdown identifies the most expensive lanes right now",
            "Average Transit Time shown alongside cost — efficiency and spend together",
            "Procurement enters negotiations with current, lane-level data as leverage",
        ],
        "outcome": "Freight cost visibility moves from a monthly finance report to a live dashboard. Procurement negotiates from a position of current, accurate data.",
        "capability": "Live Cost per Shipment & Route Analytics — TMS",
        "ai_feature": "Gen AI Assistant: Team can ask 'What worked last time freight costs spiked on this lane?' — the assistant draws on historical incident knowledge and recommends a proven course of action.",
        "steps": [
            ("Cost Spike Visible","Cost per Shipment KPI rises — route-level breakdown available live"),
            ("Lane Identified",   "Most expensive lane pinpointed without any manual data extraction"),
            ("Negotiate with Data","Procurement enters carrier talks with current, accurate lane cost data"),
        ],
    },
    {
        "num": 8,
        "layer": "Unified Dashboard — Leadership View",
        "accent": GREEN,
        "accent_light": GREEN_LIGHT,
        "title": "One View for Leadership — Across All Warehouses & Categories",
        "trigger": "The supply chain head needs a single view across warehouses in Delhi, Mumbai, and Bangalore for all product categories. Today it requires three logins and a manual consolidation.",
        "without": [
            "Three separate tools — WMS, TMS, Planning — each needing a different login",
            "No single composite metric covering the end-to-end supply chain",
            "Leadership report built manually each week — always a week behind",
            "Strategic decisions made on a fragmented, time-delayed picture",
        ],
        "with_tool": [
            "Single login — WMS, TMS, and Planning KPIs visible in one dashboard",
            "Filter by Warehouse (DEL / MUM / BLR) and Category in one control bar",
            "SC Cycle Time, Supply Chain Reliability, Carrier-to-Shelf Days — all in one view",
            "VAPI Calling Feature and Gen AI Assistant available for instant queries",
        ],
        "outcome": "Leadership gets a real-time, single-screen view of the entire supply chain — filtered to any warehouse or category — with no manual report required.",
        "capability": "Unified KPI Dashboard — WMS + TMS + Planning",
        "ai_feature": "All Three AI Features Active: Email alerts notify the right executive instantly. VAPI calls them for a recorded justification. Gen AI Assistant recommends what to do — based on everything that has happened before.",
        "steps": [
            ("Single Login",     "One URL — WMS, TMS, and Planning KPIs all visible immediately"),
            ("Filter & Focus",   "Drill into any warehouse (DEL/MUM/BLR) or category in one click"),
            ("AI on Demand",     "VAPI call or Gen AI query answers any question — no analyst needed"),
        ],
    },
]


# ══════════════════════════════════════════════════════════════════════════════
# SCENARIO SLIDE BUILDER
# ══════════════════════════════════════════════════════════════════════════════

for sc in scenarios:
    slide = prs.slides.add_slide(BLANK)
    fill_bg(slide, OFF_WHITE)

    accent  = sc["accent"]
    a_light = sc["accent_light"]

    # Left accent stripe
    add_rect(slide, 0, 0, 0.22, 7.5, fill=accent)

    # Top header band
    add_rect(slide, 0.22, 0, 13.11, 1.1, fill=NAVY)

    # Scenario label
    add_text(slide, f"SCENARIO  {sc['num']}  /  8   \u00b7   {sc['layer']}",
             0.42, 0.1, 12.5, 0.28, size=8, bold=True, color=accent)

    # Title
    add_text(slide, sc["title"],
             0.42, 0.42, 12.0, 0.58, size=19, bold=True, color=WHITE)

    # Horizontal rule below header
    add_rect(slide, 0.22, 1.1, 13.11, 0.04, fill=accent)

    # ── TRIGGER BAR ──────────────────────────────────────────────────────────
    add_rect(slide, 0.35, 1.22, 12.63, 0.62, fill=WHITE,
             line=LIGHT_GREY, line_width=Pt(1))
    add_rect(slide, 0.35, 1.22, 0.06, 0.62, fill=RED)
    add_text(slide, "SITUATION:", 0.55, 1.27, 1.4, 0.22,
             size=8, bold=True, color=RED)
    add_text(slide, sc["trigger"], 1.95, 1.27, 10.8, 0.5,
             size=9.5, color=DARK_GREY, italic=True)

    # ── WITHOUT column ───────────────────────────────────────────────────────
    add_rect(slide, 0.35, 2.0, 4.45, 3.55, fill=RED_LIGHT,
             line=RGBColor(0xFC,0xCA,0xCA), line_width=Pt(1))
    add_rect(slide, 0.35, 2.0, 4.45, 0.4, fill=RED)
    add_text(slide, "  \u274c  WITHOUT THIS PLATFORM",
             0.42, 2.04, 4.25, 0.3, size=9.5, bold=True, color=WHITE)
    y = 2.54
    for b in sc["without"]:
        add_rect(slide, 0.55, y+0.1, 0.07, 0.07, fill=RED)
        add_text(slide, b, 0.72, y, 3.95, 0.44, size=9.5, color=DARK_GREY)
        y += 0.52

    # ── WITH TOOL column ─────────────────────────────────────────────────────
    add_rect(slide, 5.0, 2.0, 4.45, 3.55, fill=GREEN_LIGHT,
             line=RGBColor(0xA7,0xF3,0xD0), line_width=Pt(1))
    add_rect(slide, 5.0, 2.0, 4.45, 0.4, fill=GREEN)
    add_text(slide, "  \u2705  WITH OUR AI PLATFORM",
             5.08, 2.04, 4.25, 0.3, size=9.5, bold=True, color=WHITE)
    y = 2.54
    for b in sc["with_tool"]:
        add_rect(slide, 5.18, y+0.1, 0.07, 0.07, fill=GREEN)
        add_text(slide, b, 5.35, y, 3.95, 0.44, size=9.5, color=DARK_GREY)
        y += 0.52

    # ── How It Works flow panel ───────────────────────────────────────────────
    flow_panel(slide, 9.65, 2.0, 3.33, 3.55, sc["steps"], accent)

    # ── Outcome bar ───────────────────────────────────────────────────────────
    add_rect(slide, 0.35, 5.72, 8.3, 0.65, fill=a_light,
             line=accent, line_width=Pt(1))
    add_rect(slide, 0.35, 5.72, 0.07, 0.65, fill=accent)
    add_text(slide, "OUTCOME", 0.55, 5.74, 1.2, 0.22,
             size=7.5, bold=True, color=accent)
    add_text(slide, sc["outcome"], 0.55, 5.95, 8.0, 0.36,
             size=9.5, color=DARK_GREY)

    # ── AI Feature strip ──────────────────────────────────────────────────────
    add_rect(slide, 0.35, 6.52, 12.63, 0.55, fill=RGBColor(0x0B, 0x1F, 0x3B))
    add_rect(slide, 0.35, 6.52, 0.07,  0.55, fill=TEAL)
    add_text(slide, "\u2728  AI FEATURE",
             0.55, 6.54, 2.0, 0.22, size=7.5, bold=True, color=TEAL)
    add_text(slide, sc["ai_feature"],
             0.55, 6.73, 12.0, 0.28, size=8.5, color=WHITE)

    # ── Metric badges ─────────────────────────────────────────────────────────
    # Single capability badge replacing the two metric badges
    add_rect(slide, 9.65, 5.72, 3.33, 0.1,  fill=sc["accent"])
    add_rect(slide, 9.65, 5.82, 3.33, 0.62, fill=WHITE,
             line=LIGHT_GREY, line_width=Pt(1))
    add_text(slide, "KEY CAPABILITY", 9.72, 5.84, 3.2, 0.2,
             size=7, bold=True, color=sc["accent"], align=PP_ALIGN.CENTER)
    add_text(slide, sc["capability"], 9.72, 6.04, 3.2, 0.35,
             size=8.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    footer(slide, left_text="AI Supply Chain Intelligence Platform", right_text=f"Scenario {sc['num']}  /  8")


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE — SUMMARY & ROI
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK)
fill_bg(slide, OFF_WHITE)
top_band(slide, NAVY, 1.05)
accent_bar(slide, color=TEAL)

add_text(slide, "WHAT YOU ACHIEVE",
         0.45, 0.14, 12, 0.3, size=8.5, bold=True, color=TEAL)
add_text(slide, "From Reactive to Proactive — Across Every Layer of Your Supply Chain",
         0.45, 0.47, 12, 0.48, size=19, bold=True, color=WHITE)

add_rect(slide, 0.22, 1.05, 13.11, 0.04, fill=TEAL)

# 6 ROI cards, 2 rows × 3 cols
roi = [
    (GREEN, "Same-Day Detection",
     "KPI breaches and exceptions surface the same day they occur — not at month-end."),
    (AMBER, "Zero Manual Reporting",
     "Live dashboards replace Excel consolidation entirely. BI lag is eliminated."),
    (TEAL,  "One Platform, One Login",
     "WDE + TMS + Planning accessible from a single URL. No tool-switching, no silos."),
    (GREEN, "Institutional Memory",
     "Every incident, root cause, and resolution is stored and queryable via Gen AI."),
    (AMBER, "Proactive Early Warnings",
     "Amber alerts fire before red crises. Act before SLAs are missed."),
    (TEAL,  "Voice-Enabled Intelligence",
     "Query any KPI out loud during a meeting. Answer in seconds, not days."),
]

for i, (color, title, desc) in enumerate(roi):
    row = i // 3
    col = i % 3
    bx = 0.35 + col * 4.3
    by = 1.22 + row * 2.08
    add_rect(slide, bx, by,      4.05, 1.82, fill=WHITE,
             line=LIGHT_GREY, line_width=Pt(1))
    add_rect(slide, bx, by,      4.05, 0.1,  fill=color)
    add_rect(slide, bx, by+0.1,  0.06, 1.72, fill=color)
    add_text(slide, title, bx+0.2, by+0.2,  3.7, 0.32, size=11, bold=True, color=NAVY)
    add_text(slide, desc,  bx+0.2, by+0.62, 3.7, 0.95, size=9.5, color=DARK_GREY)

# CTA strip
add_rect(slide, 0.35, 5.52, 12.63, 1.38, fill=NAVY)
add_rect(slide, 0.35, 5.52, 0.07,  1.38, fill=TEAL)
add_text(slide, "Ready to deploy. Live today.",
         0.6, 5.65, 12.1, 0.42, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, "baf-unified.vercel.app",
         0.6, 6.1, 12.1, 0.35, size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
add_text(slide,
         "Full AI stack  \u00b7  No infra setup required  \u00b7  Demo available right now",
         0.6, 6.5, 12.1, 0.3, size=9, color=MID_GREY, align=PP_ALIGN.CENTER)

footer(slide, right_text="Summary & Next Steps")


# ══════════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════════
output_path = r"d:\wms1\AI_Supply_Chain_Platform_Presentation.pptx"
prs.save(output_path)
print(f"Saved: {output_path}  ({len(prs.slides)} slides)")
